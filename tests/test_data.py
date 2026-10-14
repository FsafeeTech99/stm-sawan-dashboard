from io import BytesIO
import unittest

import pandas as pd
from openpyxl import load_workbook
import xlwt

from dashboard.data import (
    STM_COLUMNS, ParsedSheet, combine_records, details, excel_bytes, filter_records,
    category_summary, money, normalize_code, parse_sheet, prepare_records, read_workbook, reconcile, signed_totals, summarize,
)
from dashboard.auth import hash_password, verify_password, LoginThrottle


def workbook_bytes(headers, rows, metadata=7, legacy=False):
    output = BytesIO()
    if legacy:
        wb = xlwt.Workbook()
        ws = wb.add_sheet("รายละเอียด")
        for i in range(metadata):
            ws.write(i, 0, "รายงาน สปสช. รหัส 11450 จำนวนเงิน รายการ")
        for col, value in enumerate(headers):
            ws.write(metadata, col, value)
        for row_index, row in enumerate(rows, start=metadata + 1):
            for col, value in enumerate(row):
                if value is not None:
                    ws.write(row_index, col, value)
        wb.save(output)
    else:
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([["รายงาน สปสช. รหัส 11450 จำนวนเงิน รายการ"]] * metadata).to_excel(
                writer, index=False, header=False, sheet_name="รายละเอียด")
            pd.DataFrame(rows, columns=headers).to_excel(writer, index=False, startrow=metadata, sheet_name="รายละเอียด")
    return output.getvalue()


def records(code, amount, period="2569-09", filename="a.xlsx", extra=None):
    raw = pd.DataFrame({"HSUB": code, "งบพึงรับ": amount, "งวด": period, **(extra or {})})
    parsed = ParsedSheet(filename, "Sheet1", raw.assign(_source_file=filename, _source_sheet="Sheet1", _source_row=range(2, len(raw) + 2)), 1, 1)
    return prepare_records(parsed, "HSUB", "งบพึงรับ", period_column="งวด")[0]


class ExcelTests(unittest.TestCase):
    def test_metadata_5_to_8_and_both_formats(self):
        for legacy in (False, True):
            for offset in (5, 6, 7, 8):
                with self.subTest(legacy=legacy, offset=offset):
                    payload = workbook_bytes(["HSUB", "HCODE", "ยอดชดเชยทั้งสิ้น"], [[5542.0, 11450, "1,200.25"]], offset, legacy)
                    sheets, errors = read_workbook(payload, "stm.xls" if legacy else "stm.xlsx")
                    self.assertFalse(errors)
                    self.assertEqual(sheets[0].header_row, offset + 1)
                    self.assertEqual(sheets[0].frame.HSUB.tolist(), ["05542"])
                    mapped, audit = prepare_records(sheets[0], "HSUB", "ยอดชดเชยทั้งสิ้น")
                    self.assertEqual(summarize(mapped)["ยอดเงิน"].iloc[0], 1200.25)

    def test_code_and_amount(self):
        for value in (5542, 5542.0, "05542", "5542.0", " 05542 "):
            self.assertEqual(normalize_code(value), "05542")
        for value in (None, float("nan"), "123456", "abc", "12.3"):
            self.assertEqual(normalize_code(value), "")
        self.assertEqual(money("(1,234.56)"), -123456)
        self.assertEqual(money("0.10"), 10)
        for value in (None, "", "bad", "NaN", "Infinity"):
            self.assertIsNone(money(value))

    def test_filters_exact_code_and_literal_name(self):
        frame = records([5542, 15542, 11450], [100, 200, 300])
        for query in ("05542", "5542", "คำสะอาด"):
            self.assertEqual(filter_records(frame, query)["_facility"].tolist(), ["05542"])
        self.assertTrue(filter_records(frame, "[").empty)
        self.assertEqual(len(filter_records(frame, facilities=["11450"], periods=["2569-09"])), 1)

    def test_duplicate_column_order_and_precision(self):
        headers = ["HSUB", "HCODE", "DRUG", "DRUG", "ยอดชดเชยที่จ่ายจริง", "ยอดชดเชยที่จ่ายจริง", "ยอดชดเชยทั้งสิ้น", "AdjRW"]
        sheets, _ = read_workbook(workbook_bytes(headers, [[5542, 11450, 11, 22, 33, 44, 110, 0.123456]]), "a.xlsx")
        self.assertEqual(sheets[0].frame["DRUG__2"].iloc[0], 22)
        frame, _ = prepare_records(sheets[0], "HSUB", "ยอดชดเชยทั้งสิ้น")
        output = details(frame, "stm")
        self.assertEqual(len(output.columns), len(STM_COLUMNS))
        self.assertEqual(output["ยอดชดเชยที่จ่ายจริง (2)"].iloc[0], 44)
        self.assertEqual(output["AdjRW"].iloc[0], 0.123456)
        self.assertNotIn("HSUB", output)

    def test_receivable_template_and_provider_not_paid_hospital(self):
        headers = [c.split("__")[0] for c in STM_COLUMNS]
        rows = []
        for i, amount in enumerate([-100, -100, 0, 150, 150, 150, 150]):
            row = [0] * 43
            row[0], row[2], row[13], row[16] = "680300024", str(i), "05542", "11450"
            row[40] = amount
            row[42] = amount
            rows.append(row)
        parsed, errors = read_workbook(workbook_bytes(headers, rows, 0), "receivable.xlsx")
        self.assertFalse(errors)
        self.assertFalse(parsed[0].warnings)
        frame, _ = prepare_records(parsed[0], "HCODE", "ยอดชดเชยทั้งสิ้น", period_column="REP")
        summary = category_summary(frame, columns=["PP", "FS"])
        self.assertEqual(summary["รหัสหน่วยบริการ"].tolist(), ["05542"])
        self.assertEqual(summary["PP"].iloc[0], 400)
        self.assertEqual(summary["ยอดชดเชยทั้งสิ้น"].iloc[0], 400)
        self.assertEqual(signed_totals(frame), (600, -200))
        self.assertEqual(details(frame, "budget").columns.tolist(), details(frame, "stm").columns.tolist())
        self.assertEqual(len(details(frame, "budget").columns), 43)

    def test_category_missing_values_do_not_become_zero(self):
        frame = records([5542, 5542], [50, 25], extra={"PP": [50, None], "FS": [0, 0]})
        output = category_summary(frame, columns=["PP", "FS", "HC"])
        self.assertTrue(pd.isna(output["PP"].iloc[0]))
        self.assertTrue(pd.isna(output["HC"].iloc[0]))
        self.assertEqual(output["FS"].iloc[0], 0)

    def test_merged_headers_and_footer(self):
        raw = pd.DataFrame([
            ["รายงาน", None, None, None],
            ["HSUB", "HCODE", "OP", "OP"],
            [None, None, "ยอดชดเชยที่คำนวณได้", "ยอดชดเชยทั้งสิ้น"],
            [5542, 11450, 110, 100],
            ["รวมทั้งสิ้น", None, 110, 100],
            ["HSUB", "HCODE", "ยอดชดเชยที่คำนวณได้", "ยอดชดเชยทั้งสิ้น"],
        ])
        sheet = parse_sheet(raw, "a.xlsx", "data")
        self.assertEqual(sheet.header_row, 2)
        self.assertEqual(sheet.header_depth, 2)
        self.assertEqual(len(sheet.frame), 1)
        self.assertEqual(sheet.frame["_source_row"].iloc[0], 4)

    def test_invalid_data_excluded_and_diagnostics(self):
        frame = records([5542, "bad", 5542, 5542], ["1,000.01", 500, "invalid", ""])
        self.assertEqual(int(frame["_valid"].sum()), 1)
        self.assertEqual(summarize(frame)["ยอดเงิน"].iloc[0], 1000.01)
        self.assertTrue(frame["_issue"].iloc[2])

    def test_no_header_and_manual_header(self):
        raw = pd.DataFrame([["title", None], ["หน่วย", "สิทธิเงิน"], [5542, 100]])
        with self.assertRaises(ValueError):
            parse_sheet(raw, "file.xlsx", "sheet")
        parsed = parse_sheet(raw, "file.xlsx", "sheet", override=1)
        self.assertEqual(parsed.frame["สิทธิเงิน"].tolist(), [100])

    def test_reject_fake_excel_and_renamed_xlsx(self):
        with self.assertRaises(ValueError):
            read_workbook(b"<html>report</html>", "fake.xls")
        result, errors = read_workbook(workbook_bytes(["HSUB", "HCODE", "จำนวนเงิน"], [[5542, 11450, 5]]), "renamed.xls")
        self.assertEqual(len(result), 1)
        self.assertFalse(errors)

    def test_multiple_sheets_with_invalid_cover(self):
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([["หน้าปก"]]).to_excel(writer, index=False, header=False, sheet_name="ปก")
            for sheet in ("A", "B"):
                pd.DataFrame({"HSUB": [5542], "HCODE": [11450], "จำนวนเงิน": [5]}).to_excel(writer, index=False, sheet_name=sheet)
        sheets, errors = read_workbook(output.getvalue(), "all.xlsx")
        self.assertEqual(len(sheets), 2)
        self.assertEqual(len(errors), 1)

    def test_exact_duplicates_only_between_sources(self):
        a = records([5542, 5542], [100, 100], filename="a.xlsx")
        b = records([5542, 5542], [100, 90], filename="b.xlsx")
        combined, count = combine_records([a, b])
        self.assertEqual(count, 1)
        self.assertEqual(len(combined), 3)
        self.assertEqual(summarize(combined)["ยอดเงิน"].iloc[0], 290)
        self.assertEqual(len(combine_records([a, b], False)[0]), 4)

    def test_reconcile_period_and_unmatched(self):
        budget = records([5542, 5542, 11450], [100.01, 200, 500], period=["2569-09", "2569-09", "2569-08"])
        stm = records([5542, 11450], [250, 450], period=["2569-09", "2569-09"])
        output = reconcile(budget, stm)
        self.assertEqual(len(output), 3)
        self.assertEqual(output.loc[output["รหัสหน่วยบริการ"].eq("05542"), "ผลต่างรอตรวจสอบ"].iloc[0], 50.01)
        self.assertEqual(output["ผลต่างรอตรวจสอบ"].notna().sum(), 1)

    def test_reconcile_no_period_and_one_side(self):
        a = records([5542], [100], period="")
        output = reconcile(a, a)
        self.assertTrue(output["ผลต่างรอตรวจสอบ"].isna().all())
        self.assertEqual(output["สถานะ"].iloc[0], "ต้องระบุงวดก่อนเปรียบเทียบ")
        output = reconcile(records([5542], [100]), pd.DataFrame())
        self.assertTrue(output["ยอดชดเชย STM"].isna().all())
        self.assertTrue(output["ผลต่างรอตรวจสอบ"].isna().all())

    def test_export_codes_formula_and_no_disk(self):
        payload = excel_bytes({"รายละเอียด": pd.DataFrame({"HCODE": ["05542"], "HN": ["0001"], "ชื่อ": ["=1+1"], "ยอด": [100.01]})})
        wb = load_workbook(BytesIO(payload))
        ws = wb["รายละเอียด"]
        self.assertEqual(ws["A2"].value, "05542")
        self.assertEqual(ws["A2"].number_format, "@")
        self.assertEqual(ws["C2"].data_type, "s")
        self.assertEqual(ws.freeze_panes, "A2")
        duplicate_headers = pd.DataFrame([[10, 20]], columns=["DRUG", "DRUG"])
        ws = load_workbook(BytesIO(excel_bytes({"Reference": duplicate_headers})))["Reference"]
        self.assertEqual([c.value for c in ws[1]], ["DRUG", "DRUG"])
        self.assertEqual([c.value for c in ws[2]], [10, 20])

    def test_45_workbooks_combined(self):
        groups = []
        for count in (20, 25):
            frames = []
            for i in range(count):
                headers = ["HSUB", "HCODE", "ยอดชดเชยทั้งสิ้น", "งวด", "TRAN_ID"]
                rows = [[5542, 11450, 0.10, "2569-09", f"{count}-{i}"]]
                sheets, errors = read_workbook(workbook_bytes(headers, rows, metadata=5 + i % 4, legacy=i % 2 == 0), f"{i}.xls")
                self.assertFalse(errors)
                frames.append(prepare_records(sheets[0], "HSUB", "ยอดชดเชยทั้งสิ้น", period_column="งวด")[0])
            groups.append(combine_records(frames)[0])
        result = reconcile(*groups)
        self.assertEqual(result["งบพึงรับ"].iloc[0], 2)
        self.assertEqual(result["ยอดชดเชย STM"].iloc[0], 2.5)
        self.assertEqual(result["ผลต่างรอตรวจสอบ"].iloc[0], -0.5)


class AuthTests(unittest.TestCase):
    def test_hash_and_invalid_password(self):
        encoded = hash_password("long-test-password")
        self.assertTrue(verify_password("long-test-password", encoded))
        self.assertFalse(verify_password("wrong", encoded))
        self.assertFalse(verify_password("wrong", "invalid"))
        self.assertNotIn("long-test-password", encoded)

    def test_throttling(self):
        throttle = LoginThrottle()
        for _ in range(5):
            self.assertTrue(throttle.allowed("admin"))
            throttle.failed("admin")
        self.assertFalse(throttle.allowed("admin"))
        throttle.reset("admin")
        self.assertTrue(throttle.allowed("admin"))


if __name__ == "__main__":
    unittest.main()
