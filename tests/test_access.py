from io import BytesIO
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch
import unittest

import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from dashboard.auth import registered_users, verify_password
from dashboard.data import details, excel_bytes, restrict_to_facility
from dashboard.registry import FACILITIES, account_code, DEVELOPER_CREDIT

APP = str(Path(__file__).resolve().parents[1] / "app.py")


class AccessTests(unittest.TestCase):
    def test_exactly_19_fixed_accounts_and_the_requested_passwords(self):
        users = registered_users({"finance": {"password_hash": "ignored"}})
        self.assertEqual(len(users), 19)
        self.assertEqual(set(FACILITIES), {"05542", "05543", "05544", "05545", "05546", "05548",
                                         "05549", "05550", "05551", "05552", "05553", "05554",
                                         "05555", "05556", "05557", "05558", "05559", "05560", "13977"})
        for code in FACILITIES:
            user = users[f"Admin{code}"]
            self.assertTrue(verify_password(code, user["password_hash"]))
            self.assertEqual(user["facility_code"], code)
        self.assertNotIn("finance", users)
        self.assertNotIn("Admin11450", users)
        self.assertFalse(verify_password("5542", users["Admin05542"]["password_hash"]))
        self.assertIsNone(account_code("Admin11450"))

    def test_scope_cannot_be_changed_by_paid_code_or_prepared_filter(self):
        data = pd.DataFrame({"HSUB": [5542, 5555], "HCODE": [11450, 5542], "HCODE_PAID": [5555, 5542],
                             "_facility": ["05555", "05542"], "TRAN_ID": ["OWN", "OTHER"]})
        own = restrict_to_facility(data, "05542")
        self.assertEqual(own.TRAN_ID.tolist(), ["OWN"])
        self.assertEqual(own["_facility"].tolist(), ["05542"])
        self.assertTrue(restrict_to_facility(data, "11450").empty)
        self.assertTrue(restrict_to_facility(data.drop(columns=["HSUB", "HCODE"]), "05542").empty)
        exported = excel_bytes({"details": details(own, "stm")})
        ws = load_workbook(BytesIO(exported))["details"]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["C2"].value, "OWN")

    def test_builtin_login_filters_mixed_upload_and_locks_unit(self):
        payload = BytesIO()
        pd.DataFrame({"HSUB": [5542, 5555], "HCODE": [11450, 11450], "HCODE_PAID": [11450, 5542],
                      "REP": ["TEST", "TEST"], "TRAN_ID": ["OWN", "OTHER"],
                      "PP": [125, 900], "ยอดชดเชยทั้งสิ้น": [125, 900]}).to_excel(payload, index=False)
        upload = SimpleNamespace(name="mixed.xlsx", getvalue=lambda: payload.getvalue())
        app = AppTest.from_file(APP, default_timeout=45)
        app.secrets["auth"] = {"users": {}}
        app.run()
        self.assertFalse(app.exception)
        app.text_input[0].set_value("Admin05542")
        app.text_input[1].set_value("05542")
        with patch("streamlit.file_uploader", side_effect=lambda label, **kwargs: [upload] if label.startswith("1.") else []):
            app.button[0].click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.metric[0].value, "125.00")
            self.assertEqual(app.session_state["user"], "Admin05542")
            self.assertEqual(app.multiselect(key="facility_filter").options, ["05542 · " + FACILITIES["05542"]])
            self.assertTrue(app.multiselect(key="facility_filter").disabled)
            for table in app.dataframe:
                if "TRAN_ID" in table.value:
                    self.assertEqual(table.value.TRAN_ID.tolist(), ["OWN"])
            self.assertTrue(any(DEVELOPER_CREDIT in c.value for c in app.caption))
            app.text_input(key="search").set_value("05555").run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.metric), 0)


if __name__ == "__main__":
    unittest.main()
