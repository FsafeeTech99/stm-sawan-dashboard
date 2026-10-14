from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import hashlib
import re
import unicodedata

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from dashboard.registry import FACILITIES

FOCUS_SCHEMA_VERSION = 2


STM_COLUMNS = [
    "REP", "ลำดับที่", "TRAN_ID", "HN", "AN", "PID", "ชื่อ-นามสกุล",
    "วันเข้ารักษา", "วันจำหน่าย", "MAININSCL", "PROJCODE", "HTYPE_HCODE",
    "PROV_HCODE", "HCODE", "HTYPE_HCODE_PAID", "PROV_HCODE_PAID", "HCODE_PAID",
    "เรียกเก็บ", "พรบ.", "AdjRW", "ล่าช้า", "ล่าช้า(PS)", "CCUF", "AdjRW2",
    "อัตราจ่าย", "เงินเดือน", "จ่ายชดเชยหลังหัก พรบ.และเงินเดือน", "OP",
    "ยอดชดเชยที่คำนวณได้", "ยอดชดเชยที่จ่ายจริง", "HC", "DRUG", "AE",
    "DRUG__2", "INST", "ยอดชดเชยที่คำนวณได้__2", "ยอดชดเชยที่จ่ายจริง__2",
    "DMIS_DRUG", "Palliative care", "DMISHD", "PP", "FS", "ยอดชดเชยทั้งสิ้น",
]
DISPLAY_NAMES = {
    "DRUG__2": "DRUG (2)",
    "ยอดชดเชยที่คำนวณได้__2": "ยอดชดเชยที่คำนวณได้ (2)",
    "ยอดชดเชยที่จ่ายจริง__2": "ยอดชดเชยที่จ่ายจริง (2)",
}
CODE_COLUMNS = ["HSUB", "HCODE", "HCODE_PAID"]
IDENTIFIERS = ["REP", "TRAN_ID", "HN", "AN", "PID", "MAININSCL", "PROJCODE"]
NAME_COLUMNS = ["ชื่อหน่วยบริการ", "ชื่อ รพ.สต.", "ชื่อหน่วยงาน"]
BUDGET_AMOUNTS = ["งบพึงรับ", "ยอดพึงรับ", "จำนวนเงิน", "ยอดชดเชยทั้งสิ้น"]
STM_AMOUNTS = ["ยอดชดเชยทั้งสิ้น", "จำนวนเงิน"]
PERIOD_COLUMNS = ["REP", "งวด", "เดือน"]
DEFAULT_DIRECTORY = FACILITIES
CATEGORY_COLUMNS = [
    "จ่ายชดเชยหลังหัก พรบ.และเงินเดือน", "OP", "ยอดชดเชยที่คำนวณได้", "ยอดชดเชยที่จ่ายจริง",
    "HC", "DRUG", "AE", "DRUG__2", "INST", "ยอดชดเชยที่คำนวณได้__2",
    "ยอดชดเชยที่จ่ายจริง__2", "DMIS_DRUG", "Palliative care", "DMISHD", "PP", "FS",
]


def clean_text(value) -> str:
    if value is None or pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def token(value) -> str:
    return re.sub(r"[\s_\-./():]", "", clean_text(value)).casefold()


ALIASES = {
    "HC-DRUG": ["HC-DRUG", "HC Drug", "HC_DRUG", "ยาสมุนไพร"],
    "บริการ": ["บริการ", "ชื่อบริการ", "รายการบริการ"],
    "HSUB": ["hsub", "รหัสหน่วยบริการประจำ", "รหัส รพ.สต.", "รหัสรพสต", "hsubcode"],
    "HCODE": ["hcode", "รหัสหน่วยบริการ", "รหัสสถานพยาบาล", "รหัสโรงพยาบาล"],
    "HCODE_PAID": ["hcodepaid", "รหัสหน่วยบริการที่รับเงิน"],
    "ชื่อหน่วยบริการ": ["ชื่อหน่วยบริการ", "ชื่อสถานพยาบาล", "ชื่อโรงพยาบาล"],
    "ชื่อ รพ.สต.": ["ชื่อรพสต", "รพ.สต."],
    "ชื่อ-นามสกุล": ["ชื่อสกุล", "ชื่อ-สกุล", "ชื่อ-นามสกุล", "ชื่อ นามสกุล"],
    "งบพึงรับ": ["งบพึงรับ", "งบประมาณพึงรับ", "งบที่พึงรับ", "งบพึงได้รับ"],
    "ยอดพึงรับ": ["ยอดพึงรับ", "ยอดเงินพึงรับ", "ยอดที่พึงรับ"],
    "จำนวนเงิน": ["จำนวนเงิน", "ยอดเงิน", "amount"],
    "งวด": ["งวด", "งวดข้อมูล", "งวดจ่าย", "งวดบัญชี", "period"],
    "เดือน": ["เดือน", "เดือนปี", "เดือน/ปี"],
    "ลำดับที่": ["ลำดับ", "ลำดับที่", "no"],
    "วันเข้ารักษา": ["วันเข้ารักษา", "วันที่เข้ารักษา", "dateadm"],
    "วันจำหน่าย": ["วันจำหน่าย", "วันที่จำหน่าย", "datedsc"],
    "จ่ายชดเชยหลังหัก พรบ.และเงินเดือน": ["จ่ายชดเชยหลังหักพรบ.และเงินเดือน"],
}
LOOKUP = {token(c): c for c in STM_COLUMNS + CODE_COLUMNS + NAME_COLUMNS + BUDGET_AMOUNTS + PERIOD_COLUMNS}
LOOKUP.update({token(a): c for c, aliases in ALIASES.items() for a in aliases})


def canonical_header(value) -> str:
    text = clean_text(value)
    if not text or text.lower().startswith("unnamed"):
        return ""
    return LOOKUP.get(token(text), text)


def normalize_code(value) -> str:
    """Never turn missing/invalid values into plausible facility codes."""
    text = clean_text(value)
    if re.fullmatch(r"\d{1,5}(?:\.0+)?", text):
        return text.split(".")[0].zfill(5)
    return ""


def identifier(value) -> str:
    text = clean_text(value)
    return re.sub(r"\.0+$", "", text) if re.fullmatch(r"\d+\.0+", text) else text


def money(value) -> int | None:
    """Return exact integer satang; reject unreadable data instead of coercing to 0."""
    text = clean_text(value)
    if not text or text in {"-", "–", "—"}:
        return None
    text = text.replace(",", "").replace("฿", "").replace("บาท", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        amount = Decimal(text)
        if not amount.is_finite():
            return None
        return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except InvalidOperation:
        return None


def _header_candidate(raw: pd.DataFrame, index: int, depth: int) -> tuple[list[str], int]:
    headers = []
    for col in range(raw.shape[1]):
        cells = [canonical_header(raw.iat[r, col]) for r in range(index, index + depth)]
        # On merged headers the lower row contains the real field, but the top
        # row can hold identifiers such as HCODE while its lower cell is empty.
        known = [c for c in cells if token(c) in LOOKUP]
        headers.append(known[-1] if known else next((c for c in reversed(cells) if c), ""))
    recognised = {h for h in headers if token(h) in LOOKUP}
    has_code = bool(recognised.intersection(CODE_COLUMNS))
    has_amount = any(h in recognised for h in BUDGET_AMOUNTS + STM_AMOUNTS)
    has_record = bool(recognised.intersection({"TRAN_ID", "HN", "AN", "PID", "REP", "ลำดับที่"}))
    valid = len(recognised) >= 3 and has_code and (has_amount or has_record)
    return headers, (len(recognised) * 10 + int(has_amount) * 3 - depth if valid else -1)


def detect_header(raw: pd.DataFrame, override: int | None = None) -> tuple[int, int, list[str]]:
    if override is not None:
        if not 0 <= override < len(raw):
            raise ValueError("เลขแถวหัวตารางเกินจำนวนแถวในชีต")
        headers, _ = _header_candidate(raw, override, 1)
        return override, 1, headers
    best = (-1, 0, 1, [])
    for index in range(min(40, len(raw))):
        for depth in (1, 2, 3):
            if index + depth >= len(raw):
                continue
            # Do not swallow a first data row when extending a merged header.
            if depth > 1:
                last = raw.iloc[index + depth - 1]
                known_last = sum(token(v) in LOOKUP for v in last if clean_text(v))
                if known_last < 2:
                    continue
            headers, score = _header_candidate(raw, index, depth)
            if score > best[0]:
                best = (score, index, depth, headers)
    if best[0] < 0:
        raise ValueError("ไม่พบหัวตารางที่มีรหัสหน่วยบริการและยอดเงิน/ข้อมูลรายการ กรุณาระบุเลขแถวหัวตาราง")
    return best[1], best[2], best[3]


@dataclass
class ParsedSheet:
    filename: str
    sheet: str
    frame: pd.DataFrame
    header_row: int
    header_depth: int
    dropped_rows: int = 0
    warnings: list[str] = field(default_factory=list)
    file_digest: str = ""


def parse_sheet(raw: pd.DataFrame, filename: str, sheet: str, override=None) -> ParsedSheet:
    index, depth, headers = detect_header(raw, override)
    if override is None:
        for earlier in range(index):
            _, score = _header_candidate(raw, earlier, 1)
            if score >= 0:
                raise ValueError(f"พบหัวตารางหลายรูปแบบที่แถว {earlier + 1} และ {index + 1} กรุณาจัดคอลัมน์ให้ตรงกันหรือแยกชีตก่อนนำเข้า")
    counts: dict[str, int] = {}
    positions, names = [], []
    for pos, name in enumerate(headers):
        if not name:
            continue
        counts[name] = counts.get(name, 0) + 1
        names.append(name if counts[name] == 1 else f"{name}__{counts[name]}")
        positions.append(pos)
    frame = raw.iloc[index + depth:, positions].copy()
    frame.columns = names
    frame["_source_row"] = frame.index + 1
    data_cols = names
    blank = frame[data_cols].apply(lambda col: col.map(clean_text).eq("")).all(axis=1)
    totals = frame[data_cols].apply(lambda col: col.map(clean_text).str.match(
        r"^(?:รวม|รวมทั้งสิ้น|รวมทั้งหมด|ยอดรวม|รวมยอด|grand\s*total|total)(?:\s|$)", case=False
    )).any(axis=1)
    # Repeated headers may be removed only if their columns stay aligned.
    # A concatenated report with a different layout must not silently assign
    # later rows to the first report's financial columns.
    canonical_raw = raw.iloc[index + depth:].apply(lambda col: col.map(canonical_header))
    canonical_cells = canonical_raw.iloc[:, positions].set_axis(data_cols, axis=1)
    header_like = canonical_raw.isin(set(LOOKUP.values())).sum(axis=1).ge(3)
    for row_index in canonical_raw.index[header_like]:
        cells = canonical_raw.loc[row_index]
        recognised = set(cells).intersection(LOOKUP.values())
        if not recognised.intersection(CODE_COLUMNS) or not recognised.intersection(BUDGET_AMOUNTS + STM_AMOUNTS + IDENTIFIERS):
            continue
        shifted = [str(cells.iloc[pos]) for pos in range(len(cells)) if cells.iloc[pos] in LOOKUP.values() and cells.iloc[pos] != headers[pos]]
        if shifted:
            raise ValueError(f"หัวตารางซ้ำที่แถว {row_index + 1} เรียงคอลัมน์ต่างจากตารางแรก ({', '.join(shifted[:4])}) กรุณาจัดคอลัมน์ให้ตรงกันหรือแยกเป็นคนละชีต/ไฟล์ก่อนนำเข้า")
    repeated = canonical_cells.apply(lambda col: col.eq(col.name.split("__")[0])).sum(axis=1) >= 3
    removed = blank | totals | repeated
    frame = frame.loc[~removed].reset_index(drop=True)
    warnings = []
    for col in CODE_COLUMNS:
        if col in frame:
            original = frame[col].map(clean_text)
            frame[col] = frame[col].map(normalize_code)
            invalid = original.ne("") & frame[col].eq("")
            if invalid.any():
                warnings.append(f"{col}: รหัสไม่ถูกต้อง {invalid.sum()} รายการ (ไม่ใช้ในการรวมยอด)")
    for col in IDENTIFIERS:
        if col in frame:
            frame[col] = frame[col].map(identifier)
    frame["_source_file"] = filename
    frame["_source_sheet"] = sheet
    return ParsedSheet(filename, sheet, frame, index + 1, depth, int(removed.sum()), warnings)


def read_workbook(content: bytes, filename: str, override: int | None = None) -> tuple[list[ParsedSheet], list[str]]:
    # Use actual signature: an XLSX file may have been renamed to .xls.
    if content.startswith(b"PK"):
        engine = "openpyxl"
    elif content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        engine = "xlrd"
    else:
        raise ValueError("ไฟล์นี้ไม่ใช่ Excel .xls/.xlsx จริง หรือไฟล์เสียหาย")
    parsed, errors = [], []
    with pd.ExcelFile(BytesIO(content), engine=engine) as workbook:
        for sheet in workbook.sheet_names:
            try:
                raw = workbook.parse(sheet, header=None, dtype=object)
                if raw.dropna(how="all").empty:
                    continue
                parsed.append(parse_sheet(raw, filename, sheet, override))
            except ValueError as exc:
                errors.append(f"{filename} / {sheet}: {exc}")
    return parsed, errors


def first_column(frame: pd.DataFrame, candidates: list[str]) -> str | None:
    return next((c for c in candidates if c in frame.columns), None)


def restrict_to_facility(frame: pd.DataFrame, code: str) -> pd.DataFrame:
    """Authorize from source HSUB/HCODE, never from an editable UI filter."""
    if frame.empty:
        return frame.copy()
    column = first_column(frame, ["HSUB", "HCODE"])
    if code not in FACILITIES or column is None:
        return frame.iloc[:0].copy()
    scoped = frame.loc[frame[column].map(normalize_code).eq(code)].copy()
    if "_facility" in scoped:
        scoped["_facility"] = code
    if "_facility_name" in scoped:
        scoped["_facility_name"] = FACILITIES[code]
    return scoped


def prepare_records(parsed: ParsedSheet, facility: str, amount: str,
                    name: str | None = None, period_column: str | None = None,
                    period_override: str = "", directory=None) -> tuple[pd.DataFrame, dict]:
    frame = parsed.frame.copy()
    frame["_source_columns"] = "|" + "|".join(parsed.frame.columns) + "|"
    directory = directory or DEFAULT_DIRECTORY
    frame["_facility"] = frame[facility].map(normalize_code)
    frame["_cents"] = pd.array(frame[amount].map(money), dtype="Int64")
    frame["_period"] = (frame[period_column].map(clean_text) if period_column
                        else pd.Series("", index=frame.index))
    if period_override.strip():
        frame["_period"] = period_override.strip()
    frame["_period"] = frame["_period"].replace("", "ยังไม่ระบุงวด")
    file_names = frame[name].map(clean_text) if name else pd.Series("", index=frame.index)
    frame["_facility_name"] = [directory.get(code) or given or f"หน่วยบริการ {code}"
                               for code, given in zip(frame["_facility"], file_names)]
    invalid_code = frame["_facility"].eq("")
    invalid_money = frame["_cents"].isna()
    audit = {
        "ไฟล์": parsed.filename, "ชีต": parsed.sheet, "แถวหัวตาราง": parsed.header_row,
        "จำนวนแถวหัวตาราง": parsed.header_depth, "รายการ": len(frame),
        "ตัดแถวว่าง/รวม/หัวซ้ำ": parsed.dropped_rows,
        "คอลัมน์รหัส": facility, "คอลัมน์ยอดเงิน": amount,
        "รหัสไม่ถูกต้อง": int(invalid_code.sum()), "ยอดเงินว่าง/ไม่ถูกต้อง": int(invalid_money.sum()),
        "ไม่มีงวด": int(frame["_period"].eq("ยังไม่ระบุงวด").sum()),
    }
    frame["_valid"] = ~invalid_code & ~invalid_money
    frame["_issue"] = [" / ".join(msg for ok, msg in ((bad_code, "รหัสไม่ถูกต้อง"), (bad_money, "ยอดเงินว่าง/ไม่ถูกต้อง")) if ok)
                       for bad_code, bad_money in zip(invalid_code, invalid_money)]
    return frame, audit


def filter_records(frame: pd.DataFrame, query: str = "", facilities=None, periods=None) -> pd.DataFrame:
    if frame.empty:
        return frame.copy()
    mask = pd.Series(True, index=frame.index)
    query = query.strip()
    if query:
        code = normalize_code(query)
        mask &= (frame["_facility"].eq(code) if code else
                 frame["_facility_name"].str.contains(query, case=False, regex=False, na=False))
    if facilities:
        mask &= frame["_facility"].isin(facilities)
    if periods:
        mask &= frame["_period"].isin(periods)
    return frame.loc[mask].copy()


def combine_records(frames: list[pd.DataFrame], remove_identical=True) -> tuple[pd.DataFrame, int]:
    if not frames:
        return pd.DataFrame(), 0
    combined = pd.concat(frames, ignore_index=True)
    # Exact business-row duplicates only, including period/facility. Preserve
    # distinct corrections with the same TRAN_ID and different amounts.
    # Occurrence number keeps legitimate identical rows within a source sheet.
    business = [c for c in combined if not c.startswith("_")] + ["_facility", "_period", "_cents"]
    normalized = combined[business].apply(lambda col: col.map(clean_text))
    signatures = pd.util.hash_pandas_object(normalized, index=False).astype(str)
    occurrence = signatures.groupby([combined["_source_file"], combined["_source_sheet"], signatures]).cumcount()
    duplicate = pd.DataFrame({"signature": signatures, "occurrence": occurrence}).duplicated()
    count = int(duplicate.sum())
    return (combined.loc[~duplicate].reset_index(drop=True) if remove_identical else combined), count


def summarize(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=["รหัสหน่วยบริการ", "งวด", "ชื่อหน่วยบริการ", "รายการ", "ยอดเงิน"])
    valid = frame.loc[frame["_valid"]]
    grouped = valid.groupby(["_facility", "_period"], dropna=False).agg(
        ชื่อหน่วยบริการ=("_facility_name", "first"), รายการ=("_cents", "size"), ยอดเงิน=("_cents", "sum")
    ).reset_index().rename(columns={"_facility": "รหัสหน่วยบริการ", "_period": "งวด"})
    grouped["ยอดเงิน"] = grouped["ยอดเงิน"].map(lambda n: int(n) / 100)
    return grouped


def category_summary(frame: pd.DataFrame, split_period=False, columns=None) -> pd.DataFrame:
    selected = CATEGORY_COLUMNS if columns is None else columns
    keys = ["_facility"] + (["_period"] if split_period else [])
    output_columns = ["รหัสหน่วยบริการ", "ชื่อหน่วยบริการ"] + (["งวด / REP"] if split_period else [])
    output_columns += ["รายการ", "ยอดชดเชยทั้งสิ้น"] + [DISPLAY_NAMES.get(c, c) for c in selected]
    if frame.empty:
        return pd.DataFrame(columns=output_columns)
    result = []
    valid = frame.loc[frame["_valid"]]
    for group_keys, rows in valid.groupby(keys, dropna=False, sort=True):
        group_keys = group_keys if isinstance(group_keys, tuple) else (group_keys,)
        row = {"รหัสหน่วยบริการ": group_keys[0], "ชื่อหน่วยบริการ": rows["_facility_name"].iloc[0],
               "รายการ": len(rows), "ยอดชดเชยทั้งสิ้น": int(rows["_cents"].sum()) / 100}
        if split_period:
            row["งวด / REP"] = group_keys[1]
        for col in selected:
            label = DISPLAY_NAMES.get(col, col)
            if col not in rows:
                row[label] = None
                continue
            amounts = rows[col].map(money)
            # A blank/invalid component does not become a reassuring zero or
            # a partial category total. Complete numeric zeroes remain zero.
            row[label] = None if amounts.isna().any() else sum(int(v) for v in amounts) / 100
        result.append(row)
    return pd.DataFrame(result, columns=output_columns)


def signed_totals(frame: pd.DataFrame) -> tuple[float, float]:
    if frame.empty:
        return 0.0, 0.0
    amounts = frame.loc[frame["_valid"], "_cents"]
    return int(amounts[amounts.gt(0)].sum()) / 100, int(amounts[amounts.lt(0)].sum()) / 100


def reconcile(budget: pd.DataFrame, stm: pd.DataFrame) -> pd.DataFrame:
    left = summarize(budget).rename(columns={"ยอดเงิน": "งบพึงรับ", "รายการ": "รายการพึงรับ"})
    right = summarize(stm).rename(columns={"ยอดเงิน": "ยอดชดเชย STM", "รายการ": "รายการ STM"})
    # Join only aggregated facility+period tables; never multiply patient rows.
    joined = left.merge(right, on=["รหัสหน่วยบริการ", "งวด"], how="outer", suffixes=("_budget", "_stm"), indicator=True)
    joined["ชื่อหน่วยบริการ"] = joined["ชื่อหน่วยบริการ_budget"].combine_first(joined["ชื่อหน่วยบริการ_stm"])
    # Missing counterpart stays blank: it does not establish a zero payment.
    matched = joined["_merge"].eq("both") & joined["งวด"].ne("ยังไม่ระบุงวด")
    joined["ผลต่างรอตรวจสอบ"] = (joined["งบพึงรับ"] - joined["ยอดชดเชย STM"]).round(2).where(matched)
    joined["สถานะ"] = "รอตรวจสอบ"
    joined.loc[matched & joined["ผลต่างรอตรวจสอบ"].eq(0), "สถานะ"] = "ยอดตรงกัน"
    joined.loc[matched & joined["ผลต่างรอตรวจสอบ"].gt(0), "สถานะ"] = "พึงรับสูงกว่า STM"
    joined.loc[matched & joined["ผลต่างรอตรวจสอบ"].lt(0), "สถานะ"] = "STM สูงกว่าพึงรับ"
    joined.loc[joined["_merge"].eq("left_only"), "สถานะ"] = "ไม่มี STM งวดเดียวกัน"
    joined.loc[joined["_merge"].eq("right_only"), "สถานะ"] = "ไม่มีงบพึงรับงวดเดียวกัน"
    joined.loc[joined["งวด"].eq("ยังไม่ระบุงวด"), "สถานะ"] = "ต้องระบุงวดก่อนเปรียบเทียบ"
    cols = ["รหัสหน่วยบริการ", "ชื่อหน่วยบริการ", "งวด", "งบพึงรับ", "ยอดชดเชย STM",
            "ผลต่างรอตรวจสอบ", "รายการพึงรับ", "รายการ STM", "สถานะ"]
    return joined[cols].sort_values(["รหัสหน่วยบริการ", "งวด"]).reset_index(drop=True)


def details(frame: pd.DataFrame, group: str) -> pd.DataFrame:
    if group == "stm" or (group == "budget" and ("TRAN_ID" in frame or "ยอดชดเชยทั้งสิ้น" in frame)):
        # Full requested schema in the specified order, including absent columns.
        extra = [c for c in ["HC-DRUG", "บริการ"] if c in frame]
        result = frame.reindex(columns=STM_COLUMNS + extra).copy()
        numeric = [c for c in STM_COLUMNS if c not in IDENTIFIERS + CODE_COLUMNS + [
            "ลำดับที่", "ชื่อ-นามสกุล", "วันเข้ารักษา", "วันจำหน่าย", "HTYPE_HCODE", "PROV_HCODE",
            "HTYPE_HCODE_PAID", "PROV_HCODE_PAID"]]
        def display_number(value):
            text = clean_text(value)
            if not text:
                return None
            normalized = text.replace(",", "").replace("฿", "").replace("บาท", "").strip()
            if normalized.startswith("(") and normalized.endswith(")"):
                normalized = "-" + normalized[1:-1]
            try:
                number = Decimal(normalized)
                return float(number) if number.is_finite() else text
            except InvalidOperation:
                return text
        for col in numeric:
            # AdjRW/CCUF/rates must retain their fractional precision; flags
            # and unreadable original values also remain visible for review.
            result[col] = result[col].map(display_number)
        return result.rename(columns=DISPLAY_NAMES)
    columns = {"_facility": "HSUB", "_facility_name": "ชื่อหน่วยบริการ", "_period": "งวด",
               "_cents": "งบพึงรับ", "_source_file": "ไฟล์ต้นทาง", "_source_sheet": "ชีต", "_source_row": "แถวต้นทาง"}
    result = frame.reindex(columns=columns).rename(columns=columns)
    result["งบพึงรับ"] = result["งบพึงรับ"].map(lambda n: int(n) / 100 if pd.notna(n) else None)
    return result


def excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Memory-only export: no shared patient file on the server filesystem."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for title, frame in sheets.items():
            safe = frame.copy().astype(object).where(pd.notna(frame), None)
            # Cell-wise mapping also preserves duplicate reference header names.
            safe = safe.map(lambda v: "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v)
            safe.to_excel(writer, sheet_name=title[:31], index=False)
            ws = writer.sheets[title[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="123F4C")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 42
            for i, col in enumerate(safe.columns, start=1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(35, max(14, len(str(col)) + 3))
                is_text = col in CODE_COLUMNS + IDENTIFIERS + ["รหัสหน่วยบริการ"]
                for cells in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for cell in cells:
                        if is_text:
                            cell.number_format = "@"
                        elif isinstance(cell.value, (float, int)):
                            if col in ["AdjRW", "AdjRW2", "CCUF", "อัตราจ่าย"]:
                                cell.number_format = "0.000000"
                            elif col in ["รายการ", "รายการพึงรับ", "รายการ STM", "แถว", "แถวต้นทาง", "แถวหัวตาราง",
                                         "จำนวนแถวหัวตาราง", "ตัดแถวว่าง/รวม/หัวซ้ำ", "รหัสไม่ถูกต้อง", "ยอดเงินว่าง/ไม่ถูกต้อง", "ไม่มีงวด"]:
                                cell.number_format = "#,##0"
                            else:
                                cell.number_format = "#,##0.00;[Red](#,##0.00)"
    return output.getvalue()


def content_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
